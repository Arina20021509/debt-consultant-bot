"""Квалификация заявки: workflow + LLM.

Порядок вопросов задан кодом (workflow). LLM только понимает свободный ответ клиента
и превращает его в данные: «должна банкам около 700 тысяч» → 700000.
Предварительную оценку для юриста считает код по правилам, а не модель:
решение о банкротстве принимает юрист, бот только собирает данные.
"""

import json
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook

from . import config, i18n, llm

LEADS_PATH = config.ROOT / "data" / "leads.xlsx"

# (поле, что извлечь из ответа). Сами вопросы на разных языках лежат в i18n.QUESTIONS
FIELDS = [
    ("name", "имя клиента, строка"),
    ("debt",
     "общая сумма долгов в рублях, целое число. Примеры: «1,2 млн» → 1200000, «полтора миллиона» → 1500000, "
     "«350к» → 350000, «где-то 400» → 400000 (число меньше 1000 без единиц измерения — это тысячи рублей), "
     "по-узбекски: «700 ming» → 700000, «1,5 mln» → 1500000"),
    ("creditors", "список кредиторов через запятую, на русском языке, строка"),
    ("property", "имущество клиента кратко, на русском языке, строка; если имущества нет, верни «нет»"),
    ("contact", "телефон или ник в Telegram"),
]

CANCEL = re.compile(r"^(/cancel|отменить заявку|отмена|arizani bekor qilish|bekor qilish)$", re.IGNORECASE)

# Фразы, по которым бот понимает, что клиент хочет оставить заявку (русский и узбекский)
INTENT = re.compile(
    r"/заявка|запишите|запиши меня|записаться|оставить заявку|хочу на консультацию|перезвоните"
    r"|konsultatsiyaga yozil|yozib qo.y|ariza qoldir",
    re.IGNORECASE,
)

EXTRACT_PROMPT = """Извлеки из ответа клиента значение поля и верни ТОЛЬКО JSON вида {{"value": ...}}.
Клиент может отвечать на русском или узбекском.
Поле: {description}.
Если в ответе нет нужной информации, верни {{"value": null}}.

Вопрос: {question}
Ответ клиента: {answer}"""


def wants_to_apply(text: str) -> bool:
    return bool(INTENT.search(text))


# «к»/«k» считаем тысячами, только если стоит сразу после числа: «350к», но не «банк»
_MULTIPLIERS = [(r"млн|миллион|mln|million", 1_000_000), (r"тыс|ming|\d\s*[кk]\b", 1_000)]


def _to_int(value) -> int | None:
    """Запасной разбор суммы, если модель вернула строку: «1,2 млн» → 1200000, «700 000 ₽» → 700000."""
    if isinstance(value, (int, float)):
        return int(value)
    if not isinstance(value, str):
        return None
    match = re.search(r"\d+(?:[.,]\d+)?", re.sub(r"[\s ]", "", value))
    if not match:
        return None
    number = float(match.group(0).replace(",", "."))
    for pattern, multiplier in _MULTIPLIERS:
        if re.search(pattern, value.lower()):
            return int(number * multiplier)
    return int(number)


def parse_contact(answer: str) -> str | None:
    """Контакт проверяется кодом, а не моделью: номер телефона или @ник надёжнее найти по шаблону."""
    username = re.search(r"@[A-Za-z0-9_]{4,}", answer)
    if username:
        return username.group(0)
    if answer.lower().startswith("telegram"):
        return answer
    digits = re.sub(r"\D", "", answer)
    if len(digits) == 10 or (len(digits) == 11 and digits[0] in "78"):
        return "+7" + digits[-10:]
    if 11 <= len(digits) <= 15:
        return "+" + digits
    return None


def extract(field: str, question: str, description: str, answer: str):
    if field == "contact":
        return parse_contact(answer)
    prompt = EXTRACT_PROMPT.format(description=description, question=question, answer=answer)
    raw = llm.complete([{"role": "user", "content": prompt}], temperature=0.1)
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    try:
        value = json.loads(match.group(0))["value"] if match else None
    except (json.JSONDecodeError, KeyError):
        value = None
    return _to_int(value) if field == "debt" else value


def assess(lead: dict) -> str:
    """Предварительная оценка для юриста. Клиенту она не показывается."""
    notes = []
    debt = lead.get("debt")
    if not isinstance(debt, int):
        # Клиент дважды ответил так, что сумму не удалось разобрать, — сохраняем его слова как есть
        notes.append(f"сумма долга не распознана, ответ клиента: «{debt}»" if debt else "сумма долга не указана")
    elif debt < 25_000:
        notes.append("долг меньше 25 000 ₽, банкротство, скорее всего, нецелесообразно")
    else:
        if debt <= 1_000_000:
            notes.append("сумма подходит для внесудебного банкротства, проверить исполнительные производства на сайте ФССП")
        if debt > 500_000:
            notes.append("долг больше 500 000 ₽: при просрочке больше 3 месяцев обязан подать на судебное банкротство")
    property_text = str(lead.get("property") or "").lower()
    if re.search(r"ипотек|ipoteka", property_text):
        notes.append("есть ипотека, жильё может быть реализовано")
    if re.search(r"машин|авто|дол[яиюе]|mashina|ulush", property_text):
        notes.append("есть имущество, которое может войти в конкурсную массу")
    return "; ".join(notes)


def save_lead(lead: dict) -> None:
    header = ["Дата", "Канал", "Имя", "Сумма долга, ₽", "Кредиторы", "Имущество", "Контакт", "Оценка для юриста"]
    if LEADS_PATH.exists():
        workbook = load_workbook(LEADS_PATH)
        sheet = workbook.active
    else:
        LEADS_PATH.parent.mkdir(exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Заявки"
        sheet.append(header)
    sheet.append([
        datetime.now().strftime("%d.%m.%Y %H:%M"),
        lead.get("channel"),
        lead.get("name"), lead.get("debt"), lead.get("creditors"),
        lead.get("property"), lead.get("contact"), lead.get("assessment"),
    ])
    workbook.save(LEADS_PATH)


class Qualification:
    """Ведёт клиента по вопросам заявки. Если ответ не удалось разобрать, переспрашивает один раз."""

    lang = "ru"  # значение по умолчанию для заявок, сохранённых до появления языков

    def __init__(self, channel: str = "Консоль", lang: str = "ru"):
        self.step = 0
        self.lang = lang
        self.lead: dict = {"channel": channel}
        self.retried = False

    @property
    def current_field(self) -> str:
        return FIELDS[self.step][0]

    def _question(self, step: int) -> str:
        return i18n.question(self.lang, FIELDS[step][0])

    def first_question(self) -> str:
        return i18n.t(self.lang, "first_question", question=self._question(0))

    def handle(self, answer: str) -> tuple[str, bool]:
        """Принимает ответ клиента. Возвращает (реплика бота, заявка завершена)."""
        field, description = FIELDS[self.step]
        question = self._question(self.step)
        value = extract(field, question, description, answer)
        # Без контакта заявка бесполезна, поэтому его переспрашиваем, пока не получим
        if value is None and (not self.retried or field == "contact"):
            self.retried = True
            if field == "contact":
                return i18n.t(self.lang, "retry_contact"), False
            return i18n.t(self.lang, "retry", question=question), False
        self.lead[field] = value if value is not None else answer
        self.step, self.retried = self.step + 1, False

        if self.step < len(FIELDS):
            return self._question(self.step), False

        self.lead["assessment"] = assess(self.lead)
        save_lead(self.lead)
        name = self.lead.get("name")
        return i18n.t(self.lang, "thanks", name=f", {name}" if name else ""), True
